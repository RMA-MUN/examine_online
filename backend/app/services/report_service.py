"""考试分析报告服务：只读聚合 Task 4 输出并渲染为 xlsx，不新增表、不写新聚合。

数据收集全部复用 analytics_service / statistics_service 现有函数；
xlsx 渲染复用 score_export_service / dashboard_export_service 的 openpyxl 约定
（Workbook 多 Sheet、首行中文表头、冻结首行）。
"""

from io import BytesIO

from openpyxl import Workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.services.analytics_service import (
    get_class_compare,
    get_discrimination,
    get_exam_question_stats,
    get_exam_student_scores,
    get_knowledge_stats,
    get_score_bins,
)
from app.services.statistics_service import get_exam_statistics

REPORT_SECTIONS: tuple[str, ...] = ("scores", "quality", "knowledge")

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def normalize_report_sections(sections: list[str] | None) -> list[str]:
    """归一化报告 sections：缺省/全非法时回退全量（向后兼容）。"""
    if not sections:
        return list(REPORT_SECTIONS)
    wanted = [str(s).strip().lower() for s in sections]
    used = [s for s in REPORT_SECTIONS if s in wanted]
    return used or list(REPORT_SECTIONS)


async def collect_report_data(
    db: AsyncSession, exam_id: int, sections: list[str] | None
) -> tuple[dict, list[str]]:
    """按 sections 收集报告数据（全部调用 Task 4 / statistics 现有聚合）。"""
    used = normalize_report_sections(sections)
    data: dict = {"exam_id": exam_id}
    if "scores" in used:
        students = await get_exam_student_scores(db, exam_id, page=1, page_size=1000)
        data["scores"] = {
            "summary": await get_exam_statistics(db, exam_id),
            "bins": await get_score_bins(db, exam_id),
            "classes": await get_class_compare(db, exam_id),
            "students": students["items"],
            "total": students["total"],
        }
    if "quality" in used:
        data["quality"] = {
            "questions": await get_exam_question_stats(db, exam_id),
            "discrimination": await get_discrimination(db, exam_id),
        }
    if "knowledge" in used:
        data["knowledge"] = await get_knowledge_stats(db, exam_id)
    return data, used


def render_report_xlsx(data: dict, sections: list[str]) -> tuple[bytes, str, str]:
    """将报告数据渲染为多 Sheet xlsx（表头中文、冻结首行，与成绩导出约定一致）。"""
    exam_id = data.get("exam_id", "")
    workbook = Workbook()
    workbook.remove(workbook.active)

    if "scores" in sections and "scores" in data:
        scores = data["scores"]
        summary = scores.get("summary") or {}
        overview = workbook.create_sheet("成绩总览")
        overview.append(["指标", "数值"])
        overview.freeze_panes = "A2"
        for key, label in (
            ("total_students", "参考人数"),
            ("avg_score", "平均分"),
            ("max_score", "最高分"),
            ("min_score", "最低分"),
            ("pass_rate", "及格率(%)"),
        ):
            overview.append([label, summary.get(key)])
        overview.append([])
        overview.append(["分数段", "人数"])
        for row in scores.get("bins", []):
            overview.append([row.get("label"), row.get("count")])
        overview.append([])
        overview.append(["班级", "平均分", "及格率(%)", "人数"])
        for row in scores.get("classes", []):
            overview.append([row.get("name"), row.get("avg"), row.get("pass_rate"), row.get("count")])

        detail = workbook.create_sheet("学生成绩")
        detail.append(["排名", "姓名", "学号", "分数", "切屏次数", "状态"])
        detail.freeze_panes = "A2"
        for row in scores.get("students", []):
            student = row.get("student", {})
            detail.append(
                [row.get("rank"), student.get("name"), student.get("username"),
                 row.get("score"), row.get("switch_count"), row.get("status")]
            )

    if "quality" in sections and "quality" in data:
        quality = data["quality"]
        by_q = {row.get("question_id"): row for row in quality.get("discrimination", [])}
        sheet = workbook.create_sheet("题目质量")
        sheet.append(["题号ID", "题型", "平均分", "正确率", "难度P", "区分度D", "知识点"])
        sheet.freeze_panes = "A2"
        for row in quality.get("questions", []):
            dd = by_q.get(row.get("question_id"), {})
            sheet.append(
                [row.get("question_id"), row.get("type"), row.get("avg_score"),
                 row.get("correct_rate"), dd.get("p", row.get("p")),
                 dd.get("d", row.get("d")), row.get("knowledge")]
            )

    if "knowledge" in sections and "knowledge" in data:
        sheet = workbook.create_sheet("知识点")
        sheet.append(["知识点", "得分率(%)"])
        sheet.freeze_panes = "A2"
        for row in data["knowledge"]:
            sheet.append([row.get("name"), row.get("rate")])

    output = BytesIO()
    workbook.save(output)
    return (
        output.getvalue(),
        XLSX_MEDIA_TYPE,
        f"exam-{exam_id}-report.xlsx" if exam_id != "" else "exam-report.xlsx",
    )
